"""회귀 테스트 — 발주처 지정 출력 규격(meta.output_spec).

공고 값이 내부 프로파일보다 우선하고, 생성기와 검사기가 같은 해석을 쓰며, 모르는 값은
조용히 기본값으로 떨어지지 않아야 한다(fail-closed).
"""
from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

REPO = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))
from test_support import build_deck_cached, run_script  # noqa: E402

DOC = REPO / "skills/create-proposal-document"
BD = DOC / "scripts/build_deck.py"
DC = DOC / "scripts/deck_check.py"
FIX = DOC / "fixtures/e2e-mini-rfp/slides.json"
sys.path.insert(0, str(DOC / "scripts"))
import deck_profiles as dp  # noqa: E402

try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:  # pragma: no cover
    HAS_PPTX = False

SOURCE = "제안요청서 제출 규격 조항(가상)"


class ResolveTests(unittest.TestCase):
    def test_absent_spec_falls_back_to_profile(self):
        self.assertEqual(dp.resolve_output_spec(None, None), {})
        self.assertIs(dp.effective_style("presentation", {}), dp.get("presentation"))

    def test_source_is_required(self):
        with self.assertRaisesRegex(ValueError, "source"):
            dp.resolve_output_spec({"page_limit": 30}, None)

    def test_unknown_key_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "모르는 키"):
            dp.resolve_output_spec({"source": SOURCE, "margin_mm": 20}, None)

    def test_bundle_keys_are_redirected_not_ignored(self):
        with self.assertRaisesRegex(ValueError, "attachments"):
            dp.resolve_output_spec({"source": SOURCE, "anonymous_copy": True}, None)

    def test_unknown_canvas_is_rejected(self):
        with self.assertRaisesRegex(ValueError, "모르는 값"):
            dp.resolve_output_spec({"source": SOURCE, "canvas": "B5"}, None)

    def test_known_but_ungridded_canvas_is_routed_away(self):
        for canvas in ("A4-portrait", "A3-landscape", "4:3"):
            with self.subTest(canvas=canvas):
                with self.assertRaisesRegex(ValueError, "지정 양식 또는 DOCX"):
                    dp.resolve_output_spec({"source": SOURCE, "canvas": canvas}, None)

    def test_font_floor_above_profile_body_asks_for_bigger_profile(self):
        with self.assertRaisesRegex(ValueError, "더 큰 output_profile"):
            dp.resolve_output_spec({"source": SOURCE, "font_min_pt": 12}, "detailed-submission")
        self.assertEqual(dp.resolve_output_spec({"source": SOURCE, "font_min_pt": 12}, "presentation")
                         ["font_min_pt"], 12.0)

    def test_numbers_must_be_positive_and_page_limit_integral(self):
        for bad in ({"page_limit": 0}, {"page_limit": 12.5}, {"file_size_limit_mb": -1},
                    {"font_min_pt": True}):
            with self.subTest(bad=bad):
                with self.assertRaises(ValueError):
                    dp.resolve_output_spec({"source": SOURCE, **bad}, None)

    def test_font_floor_lifts_small_text_for_generator_and_checker(self):
        spec = dp.resolve_output_spec({"source": SOURCE, "font_min_pt": 10}, "detailed-submission")
        style = dp.effective_style("detailed-submission", spec)
        # 생성기는 표·범례를 table-1로 그린다 → 하한 이상이 되도록 table은 하한+1
        self.assertGreaterEqual(style["sizes"]["table"] - 1, 10)
        self.assertGreaterEqual(style["sizes"]["body"] - 1, 10)
        self.assertEqual(dp.get("detailed-submission")["sizes"]["table"], 10, "원본 프로파일은 불변")

    def test_stamp_round_trips(self):
        spec = dp.resolve_output_spec({"source": SOURCE, "canvas": "16:9", "page_limit": 40,
                                       "font_min_pt": 10, "file_size_limit_mb": 20}, None)
        self.assertEqual(dp.read_spec_stamp(dp.spec_stamp(spec), None), (spec, "valid"))
        self.assertEqual(dp.read_spec_stamp("", None), ({}, "missing"))
        self.assertEqual(dp.read_spec_stamp(dp.OUTPUT_SPEC_PREFIX + "{broken", None), ({}, "invalid"))


@unittest.skipUnless(HAS_PPTX, "python-pptx 없음")
class PipelineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def _spec(self, name: str, output_spec: dict, *, drop_legacy_limit: bool = True) -> Path:
        data = json.loads(FIX.read_text(encoding="utf-8"))
        if drop_legacy_limit:
            data["meta"].pop("page_limit", None)
        data["meta"]["output_spec"] = output_spec
        path = self.dir / f"{name}.json"
        path.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return path

    def _build(self, name: str, output_spec: dict, *extra: object) -> tuple[Path, object]:
        out = self.dir / f"{name}.pptx"
        return out, build_deck_cached(BD, self._spec(name, output_spec), out, *extra)

    def test_generated_deck_carries_the_spec_and_passes_its_own_check(self):
        spec = {"source": SOURCE, "canvas": "16:9", "page_limit": 40, "font_min_pt": 10}
        out, proc = self._build("ok", spec)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        stamped, state = dp.read_spec_stamp(Presentation(str(out)).core_properties.keywords, None)
        self.assertEqual(state, "valid")
        self.assertEqual(stamped["font_min_pt"], 10.0)
        check = run_script(DC, out, "--stage", "draft")
        self.assertEqual(check.returncode, 0, check.stdout)
        self.assertIn("발주처 지정 규격 적용", check.stdout)
        self.assertIn("최소 폰트 10", check.stdout)

    def test_cli_cannot_lower_the_rfp_font_floor(self):
        out, proc = self._build("floor", {"source": SOURCE, "font_min_pt": 10})
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        check = run_script(DC, out, "--min-font", "6", "--stage", "draft")
        self.assertIn("최소 폰트 10", check.stdout)

    def test_spec_page_limit_applies_without_cli_flag(self):
        out, proc = self._build("pages", {"source": SOURCE, "page_limit": 5})
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)  # non-strict: 위반만 보고
        self.assertIn("페이지 제한", proc.stdout)
        check = run_script(DC, out, "--stage", "draft")
        self.assertEqual(check.returncode, 1, check.stdout)
        self.assertIn("페이지 제한 5장", check.stdout)

    def test_conflicting_page_limits_refuse_to_build(self):
        path = self._spec("conflict", {"source": SOURCE, "page_limit": 30}, drop_legacy_limit=False)
        proc = run_script(BD, path, "-o", self.dir / "conflict.pptx")
        self.assertEqual(proc.returncode, 2)
        self.assertIn("공고 원문 값 하나로", proc.stderr)

    def test_ungridded_canvas_refuses_to_build(self):
        path = self._spec("a3", {"source": SOURCE, "canvas": "A3-landscape"})
        proc = run_script(BD, path, "-o", self.dir / "a3.pptx")
        self.assertEqual(proc.returncode, 2)
        self.assertFalse((self.dir / "a3.pptx").exists())

    def test_file_size_limit_is_checked(self):
        out, proc = self._build("size", {"source": SOURCE, "file_size_limit_mb": 0.001})
        self.assertIn("지정 한도", proc.stdout)
        check = run_script(DC, out, "--stage", "draft")
        self.assertEqual(check.returncode, 1, check.stdout)

    def test_tampered_spec_stamp_blocks_submission_warns_draft(self):
        out, proc = self._build("tamper", {"source": SOURCE, "font_min_pt": 10})
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        prs = Presentation(str(out))
        prs.core_properties.keywords = dp.OUTPUT_SPEC_PREFIX + '{"source":"x","canvas":"B5"}'
        bad = self.dir / "tampered.pptx"
        prs.save(str(bad))
        self.assertEqual(run_script(DC, bad, "--stage", "submission").returncode, 1)
        draft = run_script(DC, bad, "--stage", "draft")
        self.assertEqual(draft.returncode, 0, draft.stdout)
        self.assertIn("해석할 수 없다", draft.stdout)

    def test_no_spec_keeps_existing_profile_behaviour(self):
        out = self.dir / "plain.pptx"
        proc = build_deck_cached(BD, FIX, out)
        self.assertEqual(proc.returncode, 0, proc.stdout + proc.stderr)
        self.assertFalse(str(Presentation(str(out)).core_properties.keywords or "")
                         .startswith(dp.OUTPUT_SPEC_PREFIX))
        check = run_script(DC, out)
        self.assertNotIn("발주처 지정 규격", check.stdout)


if __name__ == "__main__":
    unittest.main()
