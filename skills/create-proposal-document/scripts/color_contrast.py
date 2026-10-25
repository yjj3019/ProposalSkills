"""Check solid RGB text contrast; do not infer theme, transparency, or image backgrounds."""
from __future__ import annotations

from pptx.enum.dml import MSO_COLOR_TYPE, MSO_FILL


def contrast_ratio(foreground: str, background: str) -> float:
    def luminance(hex6: str) -> float:
        channels = [int(hex6[i:i + 2], 16) / 255 for i in (0, 2, 4)]
        linear = [c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4
                  for c in channels]
        return sum(c * weight for c, weight in zip(linear, (0.2126, 0.7152, 0.0722)))
    light, dark = sorted((luminance(foreground), luminance(background)), reverse=True)
    return (light + 0.05) / (dark + 0.05)


def solid_rgb(fill) -> str | None:
    if fill.type != MSO_FILL.SOLID or fill.fore_color.type != MSO_COLOR_TYPE.RGB:
        return None
    # RGB transformations (alpha, tint, etc.) mean the original HEX is not the displayed color.
    if fill._xPr.xpath("./a:solidFill/a:srgbClr/*"):
        return None
    return str(fill.fore_color.rgb)


def slide_issues(slide) -> tuple[list[str], int]:
    problems: list[str] = []
    skipped = 0

    def check(frame, background, label):
        nonlocal skipped
        for paragraph in frame.paragraphs:
            for run in paragraph.runs:
                if not run.text.strip():
                    continue
                size = run.font.size or paragraph.font.size
                color = run.font.color
                if (background is None or size is None or color.type != MSO_COLOR_TYPE.RGB
                        or run._r.xpath("./a:rPr/a:solidFill/a:srgbClr/*")):
                    skipped += 1
                    continue
                bold = run.font.bold if run.font.bold is not None else paragraph.font.bold
                minimum = 3.0 if size.pt >= 18 or (size.pt >= 14 and bold) else 4.5
                measured = contrast_ratio(str(color.rgb), background)
                if measured < minimum:
                    message = f"{label}: 색상 대비 {measured:.2f}:1 < {minimum:g}:1"
                    if message not in problems:
                        problems.append(message)

    def walk(shapes, background):
        previous = []
        for shape in shapes:
            if hasattr(shape, "shapes"):
                # Compare group-local coordinates only within the same group.
                walk(shape.shapes, None)
            elif getattr(shape, "has_table", False):
                for row_no, row in enumerate(shape.table.rows, 1):
                    for col_no, cell in enumerate(row.cells, 1):
                        if not cell.is_spanned:
                            check(cell.text_frame, solid_rgb(cell.fill),
                                  f"{shape.name} 셀 {row_no},{col_no}")
            elif getattr(shape, "has_text_frame", False):
                own = solid_rgb(shape.fill)
                behind = background
                for other in reversed(previous):
                    intersects = (other.left < shape.left + shape.width and other.top < shape.top + shape.height
                                  and other.left + other.width > shape.left
                                  and other.top + other.height > shape.top)
                    if not intersects:
                        continue
                    if (other.left <= shape.left and other.top <= shape.top
                            and other.left + other.width >= shape.left + shape.width
                            and other.top + other.height >= shape.top + shape.height):
                        # Image, group, and theme fills have unknown background colors.
                        behind = solid_rgb(other.fill) if hasattr(other, "fill") else None
                    else:
                        behind = None
                    break
                effective = own or behind
                if own is None and shape.fill.type not in (None, MSO_FILL.BACKGROUND):
                    effective = None
                check(shape.text_frame, effective, shape.name)
            previous.append(shape)

    walk(slide.shapes, solid_rgb(slide.background.fill))
    return problems, skipped
