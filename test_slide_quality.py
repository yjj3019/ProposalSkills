"""외부 장표 스킬 비교에서 선정한 입력 보존·대비·노트 회귀 검사."""
from __future__ import annotations

import sys
import unittest
from pathlib import Path

from pptx.dml.color import RGBColor
from pptx.util import Pt

sys.path.insert(0, str(Path(__file__).parent / "skills/create-proposal-document/scripts"))
import build_deck
import color_contrast


class SlideQualityTests(unittest.TestCase):
    def builder(self, slide, **meta):
        return build_deck.DeckBuilder({"meta": meta, "slides": [slide]}, None, True)

    def test_table_rejects_extra_or_missing_cells(self):
        for rows in ([["a", "b", "lost"]], [["a"]], ["ab"]):
            with self.subTest(rows=rows), self.assertRaisesRegex(ValueError, "셀 수"):
                self.builder({"type": "table", "title": "T", "lead": "L",
                              "columns": ["A", "B"], "rows": rows}).build()

    def test_table_rejects_invalid_columns_widths_and_alignment(self):
        cases = [{"columns": []}, {"columns": "AB"}, {"col_widths": [1]},
                 {"col_widths": [1, 0]}, {"col_widths": [1, float("nan")]},
                 {"col_widths": [1, True]}, {"right_cols": [2]}, {"right_cols": [True]}]
        for overrides in cases:
            with self.subTest(overrides=overrides), self.assertRaises(ValueError):
                self.builder({"type": "table", "title": "T", "lead": "L",
                              "columns": ["A", "B"], "rows": [["a", "b"]], **overrides}).build()

    def test_valid_table_preserves_every_cell(self):
        builder = self.builder({"type": "table", "title": "T", "lead": "L",
                                "columns": ["A", "B"], "rows": [["a", "b"], ["c", "d"]],
                                "col_widths": [1, 2], "right_cols": [1]})
        builder.build()
        table = next(s.table for s in builder.prs.slides[0].shapes if s.has_table)
        self.assertEqual([[c.text for c in row.cells] for row in table.rows],
                         [["A", "B"], ["a", "b"], ["c", "d"]])
        self.assertEqual(builder.violations, [])

    def test_notes_survive_every_layout_and_matrix_split(self):
        import json
        fixture = Path(build_deck.__file__).parents[1] / "fixtures/e2e-mini-rfp/slides.json"
        spec = json.loads(fixture.read_text(encoding="utf-8"))
        spec["slides"].append({"type": "bullets", "title": "T", "lead": "L", "items": ["a"]})
        for slide in spec["slides"]:
            slide["notes"] = "목적: 판단 근거 설명\n전환: 다음 장에서 실행 순서를 설명"
        builder = build_deck.DeckBuilder(spec, None, True)
        builder.build()
        self.assertEqual({s["type"] for s in spec["slides"]}, build_deck.ALL_TYPES)
        for slide in builder.prs.slides:
            self.assertEqual(slide.notes_slide.notes_text_frame.text, spec["slides"][0]["notes"])

    def test_contrast_formula_and_builder_reject_unreadable_brand(self):
        self.assertAlmostEqual(color_contrast.contrast_ratio("000000", "FFFFFF"), 21)
        self.assertEqual(color_contrast.contrast_ratio("FFFFFF", "FFFFFF"), 1)
        builder = self.builder({"type": "cover"}, title="T", palette={"primary": "FFFFFF"})
        builder.build()
        self.assertTrue(any("색상 대비" in v for v in builder.violations))

    def test_actual_deck_rechecks_cover_table_and_large_text(self):
        builder = self.builder({"type": "table", "title": "T", "lead": "L",
                                "columns": ["A"], "rows": [["a"]]})
        builder.build()
        slide = builder.prs.slides[0]
        cell = next(s.table for s in slide.shapes if s.has_table).cell(0, 0)
        run = cell.text_frame.paragraphs[0].runs[0]
        cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        problems, skipped = color_contrast.slide_issues(slide)
        self.assertTrue(any("셀 1,1" in p for p in problems))
        self.assertEqual(skipped, 0)
        run.font.color.rgb = RGBColor.from_string("888888")  # 3.54:1
        run.font.size = Pt(14)
        run.font.bold = True
        self.assertEqual(color_contrast.slide_issues(slide)[0], [])
        run.font.bold = False
        self.assertTrue(color_contrast.slide_issues(slide)[0])
        run.font.size = Pt(18)
        self.assertEqual(color_contrast.slide_issues(slide)[0], [])

    def test_theme_fill_and_transparency_are_not_assumed_rgb(self):
        from pptx.enum.dml import MSO_THEME_COLOR
        from pptx.oxml.xmlchemy import OxmlElement
        builder = self.builder({"type": "table", "title": "T", "lead": "L",
                                "columns": ["A"], "rows": [["a"]]})
        builder.build()
        cell = next(s.table for s in builder.prs.slides[0].shapes if s.has_table).cell(0, 0)
        cell.fill.fore_color.theme_color = MSO_THEME_COLOR.ACCENT_1
        self.assertGreater(color_contrast.slide_issues(builder.prs.slides[0])[1], 0)
        cell.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        alpha = OxmlElement("a:alpha")
        alpha.set("val", "50000")
        cell._tc.xpath("./a:tcPr/a:solidFill/a:srgbClr")[0].append(alpha)
        self.assertGreater(color_contrast.slide_issues(builder.prs.slides[0])[1], 0)

    def test_patterned_backdrop_is_not_assumed_white(self):
        # 불명확한 채움은 같은 좌표의 흰 장표 배경으로 추정하지 않는다.
        builder = self.builder({"type": "cover"}, title="T")
        builder.build()
        slide = builder.prs.slides[0]
        band = next(s for s in slide.shapes if s.name == "COVER_BAND")
        band.fill.patterned()
        problems, skipped = color_contrast.slide_issues(slide)
        self.assertEqual(problems, [])
        self.assertGreater(skipped, 0)
        title = next(s for s in slide.shapes if s.name == "TITLE")
        title.fill.patterned()
        self.assertEqual(color_contrast.slide_issues(slide)[0], [])
        self.assertGreater(color_contrast.slide_issues(slide)[1], 0)


if __name__ == "__main__":
    unittest.main()
