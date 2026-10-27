"""Regression tests for false passes and data loss reproduced by the 2026-09 external diagnostic (F01~F05). Reuse previously passing inputs and expect blockers/usage errors; include positive controls to detect overly strict blocking."""
from __future__ import annotations

import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

REPO = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))
from test_support import Result, run_script  # noqa: E402
BEST = REPO / "skills/create-best-proposal"
DOC = REPO / "skills/create-proposal-document"
UG = BEST / "scripts/unified_gate.py"
QG = DOC / "scripts/quality_gate.py"
BAM = BEST / "scripts/build_audit_from_meta.py"
FIXTURES = BEST / "fixtures"

sys.path.insert(0, str(REPO / "skills/create-winning-proposal/scripts"))
sys.path.insert(0, str(BEST / "scripts"))
import ooxml_fixtures as fixtures  # noqa: E402
import proposal_gate as pg  # noqa: E402
import build_audit_from_meta as bam  # noqa: E402


def run(*args: object, **kw) -> Result:
    """Run skill scripts in-process by default (see test_support); env promotes execution to a child process for encoding-contract checks."""
    return run_script(Path(str(args[0])), *args[1:], **kw)


LEDGER_TEXT = "총 사업비 37억원 · 구축비 25억원 · 유지보수비(3년) 12억원"


def make_pptx(path: Path, text: str = "정상 문서 " + LEDGER_TEXT) -> Path:
    """Valid package: never use an unopenable file as a positive control. Since the unified gate compares ledger values with documents, positive controls must contain the audit amounts; mismatches would defeat the comparison."""
    return fixtures.submission_pptx(path, text)


def digest_of(path: Path) -> str:
    return "sha256:" + hashlib.sha256(path.read_bytes()).hexdigest()


def ready_audit() -> dict:
    return json.loads((FIXTURES / "audit_ready_financial.json").read_text(encoding="utf-8"))


class SubmissionValidationRegressionTests(unittest.TestCase):
    """Verify diagnosed unit, physical-artifact, and inspection-scope errors alongside positive controls."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)
        sys.path.insert(0, str(DOC / "scripts"))
        import check_numbers
        import unified_gate
        self.numbers, self.gate = check_numbers, unified_gate

    def test_uninspected_font_render_cannot_clear_submission(self):
        data = ready_audit()
        data["render"].update(verified=False, layout_checked=False, render_succeeded=True)
        failures = pg.evaluate(data)
        self.assertIn("render verification is missing or failed", failures)
        self.assertTrue(any("layout check is missing" in item for item in failures))

    def test_currency_scale_is_part_of_the_value(self):
        for value, text, matches in ((37, "총액 37억원", False),
                                     (37, "총액 37 천원", False),
                                     (3700000000, "총액 37 억원", True),
                                     (3700000, "총액 3.7백만원", True),
                                     (375000000, "총액 3.75억원", True),
                                     (37000, "총액 37천원", True),
                                     (3700000000, "총액 37억달러", False),
                                     (0.123, "총액 0.12원", False)):
            with self.subTest(value=value, text=text):
                _, checks = self.numbers.compare([{"id": "N", "value": value,
                                                   "unit": "KRW", "label": "총액"}], text)
                self.assertEqual(checks[0]["matched"], matches, checks)

    def test_label_matching_detects_swapped_and_inconsistent_values(self):
        entries = [{"id": "F", "label": "최초 응답", "value": 8,
                    "unit": "시간", "match_label": True},
                   {"id": "C", "label": "지속 응답", "value": 12,
                    "unit": "시간", "match_label": True}]
        _, checks = self.numbers.compare(entries, "최초 응답 12시간, 지속 응답 8시간")
        self.assertTrue(all(not c["matched"] for c in checks))
        _, checks = self.numbers.compare(entries, "최초 응답 8시간, 지속 응답 12시간")
        self.assertTrue(all(c["matched"] for c in checks))
        _, checks = self.numbers.compare(entries[:1], "최초 응답 8시간\n최초 응답 12시간")
        self.assertFalse(checks[0]["matched"])
        _, checks = self.numbers.compare(entries[:1], "최초 응답 8시간\n최초 응답: 지속 응답 8시간")
        self.assertFalse(checks[0]["matched"])

    def test_explicit_time_units_cannot_match_a_different_duration(self):
        for unit, text, expected in (
            ("분", "응답 15분", True), ("초", "응답 15초", True),
            ("분", "응답 15시간", False), ("초", "응답 15분", False),
            ("시간", "응답 15초", False), ("분", "응답 15분기", False),
            ("영업일", "응답 15영업일", True), ("영업일", "응답 15일", False),
            ("일", "응답 15영업일", False),
            ("시간", "응답 15 분", False), ("분", "응답 15 시간", False),
            ("영업일", "응답 15 일", False), ("일", "응답 15 영업일", False),
            ("분", "응답 15 분", True),
            ("시간", "응답 15", True), ("시간", "응답 15 | 기준", True),
            ("분", "응답 15 분석 대상", True), ("대", "대상 15 VM", True),
            ("KRW", "총액 15krw", True),
        ):
            with self.subTest(unit=unit, text=text):
                _, checks = self.numbers.compare(
                    [{"id": "N", "label": "응답", "value": 15, "unit": unit}], text)
                self.assertEqual(checks[0]["matched"], expected, checks)

    def test_nonfinite_tolerance_cannot_disable_sum_or_ratio_validation(self):
        for tolerance in (float("inf"), float("-inf"), float("nan")):
            for kind in ("sum", "ratio"):
                with self.subTest(tolerance=tolerance, kind=kind):
                    entries = [{"id": "A", "label": "기준", "value": 100, "unit": "개"},
                               {"id": "B", "label": "추가", "value": 200, "unit": "개"}]
                    bad = {"id": "T", "label": "계산값", "value": 900,
                           "unit": "개", "components": ["A", "B"], "tolerance": tolerance}
                    if kind == "ratio":
                        bad.update(unit="%", percent_of="A", amount=10)
                        bad.pop("components")
                    failures = pg.check_numbers(entries + [bad])
                    self.assertTrue(any("tolerance" in f for f in failures), failures)
                    self.assertTrue(any("합계가 맞지 않는다" in f or "비율이 맞지 않는다" in f
                                        for f in failures), failures)
        normal = [{"id": "A", "label": "기준", "value": 100, "unit": "개"},
                  {"id": "T", "label": "합계", "value": 100, "unit": "개",
                   "components": ["A"], "tolerance": 0}]
        self.assertEqual(pg.check_numbers(normal), [])

    def test_calendar_dependent_units_are_not_assumed_to_be_fixed_hours(self):
        for left, right in (("영업일", "일"), ("개월", "일"), ("월", "시간"),
                            ("년", "일"), ("개월", "년")):
            with self.subTest(left=left, right=right):
                entries = [{"id": "A", "label": "기간A", "value": 1,
                            "unit": left, "at_most": "B"},
                           {"id": "B", "label": "기간B", "value": 10000, "unit": right}]
                self.assertTrue(any("cannot be compared" in f for f in pg.check_numbers(entries)))
                entries[1]["unit"] = left
                self.assertEqual(pg.check_numbers(entries), [])
                entries[0]["value"] = 10001
                self.assertTrue(any("이하" in f for f in pg.check_numbers(entries)))
        fixed = [{"id": "A", "label": "기간A", "value": 60, "unit": "분", "at_most": "B"},
                 {"id": "B", "label": "기간B", "value": 1, "unit": "시간"}]
        self.assertEqual(pg.check_numbers(fixed), [])

    def test_label_matching_input_is_validated(self):
        entry = {"id": "N", "value": 8, "unit": "시간", "label": "응답", "match_label": "yes"}
        self.assertTrue(pg.check_numbers([entry]))
        items, _ = self.numbers.compare([entry], "응답 8시간")
        self.assertTrue(any(i.startswith("[차단]") for i in items))

    def _submission(self):
        doc = make_pptx(self.dir / "final.pptx")
        data = ready_audit()
        data["attachments"] = []
        return doc, data

    def _run(self, doc, data, *extra):
        data["render"]["artifact_hash"] = data["package"]["artifact_hash"] = digest_of(doc)
        audit = self.dir / "audit.json"
        audit.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return run(UG, audit, "--doc", doc, "--no-explain", *extra)

    def test_submission_rechecks_minute_unit_in_actual_pptx(self):
        doc, data = self._submission()
        data["numbers"].append({"id": "SLA", "label": "장애응답", "value": 15,
                                "unit": "분", "match_label": True})
        for shown, expected in (("15시간", 1), ("15분", 0)):
            with self.subTest(shown=shown):
                make_pptx(doc, "정상 문서 " + LEDGER_TEXT + " · 장애응답 " + shown)
                result = self._run(doc, data)
                self.assertEqual(result.returncode, expected, result.stdout)
                if expected:
                    self.assertIn("[차단] SLA", result.stdout)
                else:
                    self.assertIn("SUBMISSION-READY", result.stdout)

    def test_recorded_layout_approval_does_not_hide_tiny_text(self):
        from pptx import Presentation
        from pptx.util import Pt
        doc, data = self._submission()
        prs = Presentation(str(doc))
        body = next(s for s in prs.slides[0].shapes if s.name == "BODY")
        body.text_frame.paragraphs[0].runs[0].font.size = Pt(1)
        prs.save(str(doc))
        result = self._run(doc, data)
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("최소 폰트", result.stdout)
        self.assertNotIn("STATUS: SUBMISSION-READY", result.stdout)

    def test_submission_requires_a_bundle_when_attachments_are_declared(self):
        doc, data = self._submission()
        attachment = self.dir / "attachment.txt"
        attachment.write_text("synthetic evidence", encoding="utf-8")
        data["attachments"] = [{"name": attachment.name, "required": True,
                                "present": True, "sha256": digest_of(attachment)}]
        missing = self._run(doc, data)
        self.assertEqual(missing.returncode, 1, missing.stdout)
        self.assertIn("--bundle", missing.stdout)
        valid = self._run(doc, data, "--bundle", self.dir)
        self.assertEqual(valid.returncode, 0, valid.stdout)
        attachment.unlink()
        absent = self._run(doc, data, "--bundle", self.dir)
        self.assertEqual(absent.returncode, 1, absent.stdout)
        self.assertIn("첨부 파일 없음", absent.stdout)

    def test_recorded_layout_approval_does_not_hide_low_contrast(self):
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        doc, data = self._submission()
        prs = Presentation(str(doc))
        body = next(s for s in prs.slides[0].shapes if s.name == "BODY")
        body.fill.solid()
        body.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        body.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor.from_string("FFFFFF")
        prs.save(str(doc))
        result = self._run(doc, data)
        self.assertEqual(result.returncode, 1, result.stdout)
        self.assertIn("색상 대비", result.stdout)
        self.assertNotIn("STATUS: SUBMISSION-READY", result.stdout)

    def test_role_free_missing_digest_never_skips_file_checks(self):
        root = self.dir / "bundle"
        root.mkdir()
        data = {"mode": "submission", "attachments": [{"name": "missing.txt", "present": True}]}
        self.assertTrue(any("without a sha256" in f for f in pg.check_attachments(data)))
        self.assertTrue(any("첨부 파일 없음" in f for f in self.gate.verify_bundle(data, root)))
        (root / "missing.txt").write_text("synthetic evidence", encoding="utf-8")
        self.assertTrue(any("sha256" in f for f in self.gate.verify_bundle(data, root)))

    def test_bundle_paths_cannot_escape_the_root(self):
        root = self.dir / "bundle"
        root.mkdir()
        outside = self.dir / "outside.txt"
        outside.write_text("synthetic evidence", encoding="utf-8")
        for filename in ("../outside.txt", str(outside)):
            with self.subTest(filename=filename):
                data = {"attachments": [{"name": "outside.txt", "file": filename,
                                          "present": True, "sha256": digest_of(outside)}]}
                self.assertTrue(any("폴더 밖" in f for f in self.gate.verify_bundle(data, root)))

        data = {"attachments": [{"name": "invalid\x00.txt", "present": True}]}
        self.assertTrue(any("경로를 해석" in f for f in self.gate.verify_bundle(data, root)))


    def test_bundle_symlinks_respect_the_resolved_root(self):
        root = self.dir / "bundle"
        root.mkdir()
        inside = root / "inside.txt"
        outside = self.dir / "outside.txt"
        inside.write_text("synthetic inside", encoding="utf-8")
        outside.write_text("synthetic outside", encoding="utf-8")
        link = root / "linked.txt"
        try:
            link.symlink_to(inside)
        except (OSError, NotImplementedError) as exc:
            self.skipTest(f"file symlinks unavailable: {exc}")
        data = {"attachments": [{"name": link.name, "present": True,
                                  "sha256": digest_of(inside)}]}
        self.assertEqual(self.gate.verify_bundle(data, root), [])
        link.unlink()
        link.symlink_to(outside)
        data["attachments"][0]["sha256"] = digest_of(outside)
        self.assertTrue(any("폴더 밖" in p for p in self.gate.verify_bundle(data, root)))
        link.unlink()
        link.symlink_to(inside)
        data["attachments"][0]["sha256"] = digest_of(inside)
        self.assertEqual(self.gate.verify_bundle(data, root), [])
        inside.unlink()
        problems = self.gate.verify_bundle(data, root)
        self.assertTrue(problems, problems)


class SubmissionLifecycleTests(unittest.TestCase):
    """Real CLI retries must validate the current document, approval and bundle."""

    def test_repair_requires_current_numbers_approval_and_attachment(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "제안 검증"
            root.mkdir()
            doc = make_pptx(root / "최종 제안.pptx")
            attachment = root / "첨부 근거.txt"
            attachment.write_text("synthetic evidence", encoding="utf-8")
            audit = root / "검토 기록.json"
            data = ready_audit()
            data["attachments"] = [{"name": attachment.name, "required": True,
                                    "present": True, "sha256": digest_of(attachment)}]

            def check(expected, status):
                audit.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
                result = subprocess.run(
                    [sys.executable, "-I", str(UG), str(audit), "--doc", str(doc),
                     "--bundle", str(root), "--no-explain"],
                    cwd=root, capture_output=True, text=True, encoding="utf-8",
                    errors="replace", timeout=60)
                self.assertEqual(result.returncode, expected, result.stdout + result.stderr)
                self.assertIn(f"STATUS: {status}", result.stdout)
                if expected:
                    self.assertNotIn("STATUS: SUBMISSION-READY", result.stdout)

            def bind():
                data["render"]["artifact_hash"] = data["package"]["artifact_hash"] = digest_of(doc)

            bind()
            check(0, "SUBMISSION-READY")
            make_pptx(doc, "정상 문서 " + LEDGER_TEXT.replace("37억원", "40억원"))
            bind()  # A new hash alone cannot approve inconsistent financial content.
            check(1, "BLOCKED")
            make_pptx(doc)
            bind()
            data["render"]["visual_review_approved"] = False
            check(1, "BLOCKED")
            data["render"]["visual_review_approved"] = True
            check(0, "SUBMISSION-READY")
            attachment.write_text("revised synthetic evidence", encoding="utf-8")
            check(1, "BLOCKED")
            data["attachments"][0]["sha256"] = digest_of(attachment)
            check(0, "SUBMISSION-READY")


class SyntheticProposalLayoutTests(unittest.TestCase):
    """Layout regressions observed in a natively rendered synthetic proposal."""

    def setUp(self):
        sys.path.insert(0, str(DOC / "scripts"))

    def _build(self, req_ids, caption="합성 검수계획 · 시험 미실행"):
        import build_deck
        builder = build_deck.DeckBuilder({
            "meta": {"output_profile": "presentation", "output_spec": {
                "source": "합성 RFP", "font_min_pt": 14}},
            "slides": [{"type": "gantt", "title": "수행 일정", "lead": "월별 검수계획입니다.",
                        "caption": caption, "source": "합성 RFP",
                        "req_ids": req_ids, "months": 12,
                        "tasks": [{"name": "분기 검수", "start": 1, "end": 12, "phase": True,
                                   "milestones": [{"at": 6, "label": "검수"}]}]}]}, None, True)
        builder.build()
        return builder, {s.name: s for s in builder.prs.slides[0].shapes}

    def test_short_requirement_ids_keep_the_existing_column(self):
        import build_deck
        _, shapes = self._build(["R01"])
        self.assertEqual(shapes["REQID"].width, int(3.2 * build_deck.IN))

    def test_long_requirement_ids_expand_without_shrinking_text(self):
        import build_deck
        req_ids = ["R01", "R07", "R09", "R13", "R14", "R17", "R19"]
        builder, shapes = self._build(req_ids)
        req, caption = shapes["REQID"], shapes["CAPTION"]
        self.assertGreater(req.width, int(3.2 * build_deck.IN))
        self.assertLessEqual(caption.left + caption.width, req.left)
        self.assertLessEqual(req.top + req.height, shapes["PAGENO"].top)
        self.assertTrue(all(r in req.text for r in req_ids))
        self.assertEqual({r.font.size.pt for p in req.text_frame.paragraphs for r in p.runs},
                         {builder.size["caption"]})
        self.assertGreaterEqual(builder.size["caption"], 14)
        self.assertEqual(builder.violations, [])

    def test_unrenderable_requirement_ids_are_violations(self):
        for req_ids in (["R" + "9" * 300], ["R01\nR02"]):
            with self.subTest(req_ids=req_ids):
                builder, shapes = self._build(req_ids)
                self.assertTrue(any("하단 요구 ID" in v for v in builder.violations), builder.violations)
                self.assertGreater(shapes["CAPTION"].width, 0)
                self.assertGreater(shapes["REQID"].width, 0)

    def test_requirement_expansion_cannot_displace_a_fitting_caption(self):
        caption = "합성검수조건" * 5
        original, _ = self._build(["R01"], caption)
        self.assertEqual(original.violations, [])
        expanded, _ = self._build(["R01", "R07", "R09", "R13", "R14", "R17", "R19"], caption)
        self.assertTrue(any("하단 캡션" in v for v in expanded.violations), expanded.violations)

    def test_milestone_labels_have_a_checkable_contrasting_background(self):
        import color_contrast
        builder, shapes = self._build(["R01"])
        label = shapes["BODY_GMSL1"]
        self.assertEqual(color_contrast.solid_rgb(label.fill), "FFFFFF")
        run = label.text_frame.paragraphs[0].runs[0]
        self.assertGreaterEqual(color_contrast.contrast_ratio(str(run.font.color.rgb), "FFFFFF"), 4.5)
        problems, skipped = color_contrast.slide_issues(builder.prs.slides[0])
        self.assertEqual((problems, skipped), ([], 0))


class ArtifactBindingTests(unittest.TestCase):
    """F01: identify which file each verification record concerns."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def _gate(self, data: dict, *extra: str) -> subprocess.CompletedProcess:
        p = self.dir / "audit.json"
        p.write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")
        return run(UG, p, "--no-explain", *extra)

    def test_submission_without_document_is_blocked(self):
        proc = self._gate(ready_audit())
        self.assertEqual(proc.returncode, 1)
        self.assertIn("--doc", proc.stdout)

    def test_audit_only_never_claims_submission_ready(self):
        proc = self._gate(ready_audit(), "--audit-only")
        self.assertEqual(proc.returncode, 0, proc.stdout)
        self.assertIn("STATUS: AUDIT-VALID", proc.stdout)
        self.assertNotIn("STATUS: SUBMISSION-READY", proc.stdout)

    def test_matching_document_is_submission_ready(self):
        doc = make_pptx(self.dir / "final.pptx")
        data = ready_audit()
        data["attachments"] = []  # Positive control submitting only the main file.
        data["render"]["artifact_hash"] = data["package"]["artifact_hash"] = digest_of(doc)
        proc = self._gate(data, "--doc", str(doc))
        self.assertEqual(proc.returncode, 0, proc.stdout)
        self.assertIn("STATUS: SUBMISSION-READY", proc.stdout)

    def test_changed_document_cannot_reuse_the_verdict(self):
        doc = make_pptx(self.dir / "final.pptx")
        data = ready_audit()
        data["render"]["artifact_hash"] = data["package"]["artifact_hash"] = digest_of(doc)
        changed = make_pptx(self.dir / "changed.pptx", "가격이 바뀐 문서 " + LEDGER_TEXT)
        proc = self._gate(data, "--doc", str(changed))
        self.assertEqual(proc.returncode, 1)
        self.assertIn("전달된 문서와 다르다", proc.stdout)

    def test_render_and_package_must_inspect_the_same_file(self):
        data = ready_audit()
        data["package"]["artifact_hash"] = "sha256:" + "b" * 64
        failures = pg.evaluate(data)
        self.assertIn("render and package artifact_hash differ — "
                      "두 검사가 서로 다른 파일을 대상으로 했다", failures)

    def test_placeholder_hash_is_not_a_digest(self):
        data = ready_audit()
        data["render"]["artifact_hash"] = "sha256:financial-proposal-fixture"
        self.assertTrue(any("must be a sha256 digest" in f for f in pg.evaluate(data)))


class InFlightInputBindingTests(unittest.TestCase):
    """A passing child must not authorize files that changed during verification."""

    def setUp(self):
        import contextlib
        from unittest.mock import patch
        import unified_gate
        self.gate, self.patch = unified_gate, patch
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)
        self.doc = make_pptx(self.root / "candidate.pptx")
        data = ready_audit()
        data["attachments"] = []
        data["render"]["artifact_hash"] = data["package"]["artifact_hash"] = digest_of(self.doc)
        self.audit = self.root / "audit.json"
        self.audit.write_text(json.dumps(data), encoding="utf-8")
        self.original = {p: p.read_bytes() for p in (self.audit, self.doc)}
        self.stack = contextlib.ExitStack()
        self.addCleanup(self.stack.close)
        self.checks = {name: self.stack.enter_context(patch.object(self.gate, name, return_value=(0, "ok")))
                       for name in ("run_quality_gate", "run_deck_check", "run_check_numbers")}

    def invoke(self, *extra, document=True):
        import contextlib
        import io
        args = [str(self.audit), "--no-explain", *extra]
        args += ["--doc", str(self.doc)] if document else ["--audit-only"]
        output = io.StringIO()
        with contextlib.redirect_stdout(output), contextlib.redirect_stderr(output):
            code = self.gate.main(args)
        return code, output.getvalue()

    def test_changed_or_deleted_inputs_are_invalid_for_each_child(self):
        for helper in self.checks:
            for target in self.original:
                for delete in (False, True):
                    for child_code in ((0, 1) if helper == "run_quality_gate" else (0,)):
                        with self.subTest(helper=helper, target=target.name, delete=delete, code=child_code):
                            for p, content in self.original.items():
                                p.write_bytes(content)
                            for check in self.checks.values():
                                check.side_effect = None
                            def change(*args):
                                if delete:
                                    target.unlink()
                                else:
                                    target.write_bytes(target.read_bytes() + b"changed")
                                return child_code, "synthetic check"
                            self.checks[helper].side_effect = change
                            code, output = self.invoke()
                            self.assertEqual(code, 2, output)
                            self.assertIn("INVALID", output)
                            self.assertNotIn("STATUS: SUBMISSION-READY", output)

    def test_unchanged_inputs_still_pass_and_quality_failure_stays_blocked(self):
        code, output = self.invoke()
        self.assertEqual(code, 0, output)
        self.assertIn("STATUS: SUBMISSION-READY", output)
        self.checks["run_quality_gate"].return_value = (1, "quality failure")
        code, output = self.invoke()
        self.assertEqual(code, 1, output)
        self.assertIn("BLOCKED", output)

    def test_audit_only_detects_changes_during_evaluation(self):
        evaluate = pg.evaluate
        def change(data):
            failures = evaluate(data)
            self.audit.write_bytes(self.audit.read_bytes() + b"changed")
            return failures
        with self.patch.object(self.gate, "_load_gate_module", return_value=pg), \
                self.patch.object(pg, "evaluate", side_effect=change):
            code, output = self.invoke(document=False)
        self.assertEqual(code, 2, output)
        self.assertNotIn("STATUS: AUDIT-VALID", output)

    def test_bundle_is_checked_after_evaluation(self):
        attachment = make_pptx(self.root / "appendix.pptx")
        data = json.loads(self.audit.read_text(encoding="utf-8"))
        data["attachments"] = [{"name": attachment.name, "present": True, "role": "proposal",
                                "sha256": digest_of(attachment), "price_screened": True}]
        self.audit.write_text(json.dumps(data), encoding="utf-8")
        code, output = self.invoke("--bundle", str(self.root))
        self.assertEqual(code, 0, output)
        evaluate = pg.evaluate
        def change(data):
            failures = evaluate(data)
            attachment.write_bytes(attachment.read_bytes() + b"changed")
            return failures
        with self.patch.object(self.gate, "_load_gate_module", return_value=pg), \
                self.patch.object(pg, "evaluate", side_effect=change):
            code, output = self.invoke("--bundle", str(self.root))
        self.assertEqual(code, 1, output)
        self.assertIn("bundle:", output)
        self.assertNotIn("STATUS: SUBMISSION-READY", output)


class StageAndLabelTests(unittest.TestCase):
    """F02: stage, label, and explanation share one decision."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def test_submission_audit_rejects_draft_stage(self):
        proc = run(UG, FIXTURES / "audit_ready_financial.json", "--stage", "draft", "--no-explain")
        self.assertEqual(proc.returncode, 2)
        self.assertIn("--stage submission", proc.stdout + proc.stderr)

    def test_simulation_only_cannot_clear_submission(self):
        data = ready_audit()
        data["artifact_mode"] = "simulation-only"
        self.assertTrue(any("simulation-only" in f for f in pg.evaluate(data)))

    def test_draft_explain_does_not_say_submittable(self):
        data = ready_audit()
        data.update(mode="draft", artifact_required=False)
        data["checks"]["submission"] = False
        data["package"]["required"] = False
        data["submission"]["cleared"] = False
        failures = pg.evaluate(data)
        self.assertEqual(failures, [], failures)
        label, code = pg.readiness(data, failures)
        self.assertEqual((label, code), ("DRAFT-READY", 0))
        explain = pg.explain_markdown(data, [], failures)
        self.assertIn("DRAFT-READY", explain)
        self.assertNotIn("제출 가능", explain)

    def test_submission_explain_matches_its_label(self):
        data = ready_audit()
        self.assertEqual(pg.readiness(data, [])[0], "SUBMISSION-READY")
        self.assertIn("SUBMISSION-READY", pg.explain_markdown(data, [], []))


class MetaIntegrityTests(unittest.TestCase):
    """F03: conversion must not silently alter or lose inputs."""

    def test_maintenance_scope_heading_accepts_spacing_but_requires_the_section(self):
        data = {"mode": "draft", "proposal_archetype": "maintenance",
                "context": {"engagement": "operate"},
                "sections": ["지원 범위와 서비스 경계", "지원 체계", "SLA",
                             "장애 대응", "시험·검수", "인수인계"]}
        self.assertEqual(pg.check_outline(data), [])
        data["sections"][0] = "지원범위와 서비스 경계"
        self.assertEqual(pg.check_outline(data), [])
        data["sections"].pop(0)
        self.assertTrue(any("지원 범위·대상" in f for f in pg.check_outline(data)))

    def _meta(self) -> dict:
        return json.loads((FIXTURES / "meta_sample.json").read_text(encoding="utf-8"))

    def test_string_evidence_is_not_split_into_characters(self):
        meta = self._meta()
        meta.setdefault("render", {})["evidence"] = "TBD"
        with self.assertRaises(ValueError) as ctx:
            bam.build_audit(meta)
        self.assertIn("must be an array of strings", str(ctx.exception))

    def test_non_object_requirement_is_not_dropped(self):
        meta = self._meta()
        meta["requirements"] = list(meta.get("requirements", [])) + ["필수 요구 문자열"]
        with self.assertRaises(ValueError) as ctx:
            bam.build_audit(meta)
        self.assertIn("requirements[", str(ctx.exception))

    def test_requirements_need_unique_ids(self):
        data = ready_audit()
        for item in data["requirements"]:
            item.pop("id", None)
        self.assertTrue(any("lacks a non-empty id" in f for f in pg.validate_schema(data)))
        data = ready_audit()
        data["requirements"][1]["id"] = data["requirements"][0]["id"]
        self.assertTrue(any("duplicate id" in f for f in pg.validate_schema(data)))

    def test_unhashable_enum_is_usage_error_not_traceback(self):
        with tempfile.TemporaryDirectory() as tmp:
            data = ready_audit()
            data["mode"] = []
            p = Path(tmp) / "a.json"
            p.write_text(json.dumps(data), encoding="utf-8")
            proc = run(UG, p, "--no-explain")
            self.assertEqual(proc.returncode, 2)
            self.assertNotIn("Traceback", proc.stderr)
            self.assertIn("unsupported mode", proc.stdout + proc.stderr)


class EvidenceSeparationTests(unittest.TestCase):
    """F04: review status, compliance status, and evidence cannot substitute for one another."""

    def test_supported_claim_needs_evidence_in_submission(self):
        data = ready_audit()
        data["claims"][0]["evidence_refs"] = []
        self.assertTrue(any("without evidence_refs" in f for f in pg.evaluate(data)))

    def test_approved_but_unsupported_requirement_is_blocked(self):
        data = ready_audit()
        data["requirements"][0].update(support="X", fit="GAP")
        self.assertTrue(any("approved but not met" in f for f in pg.evaluate(data)))

    def test_buyer_granted_exception_is_accepted(self):
        data = ready_audit()
        data["requirements"][0].update(
            support="X", fit="GAP",
            exception={"granted_by": "발주처 계약담당", "evidence": ["질의응답 회신 2026-08-14 3항"]})
        self.assertEqual(pg.evaluate(data), [])

    def test_response_location_is_not_evidence(self):
        import bulk_matrix as bm
        rows = [bm.normalize({"id": "R1", "item": "요구", "mandatory": "필수",
                              "support": "O", "response_loc": "slide:99", "note": ""}, 1)]
        reqs = bm.to_audit_requirements(rows)
        self.assertEqual(reqs[0]["response_refs"], ["slide:99"])
        self.assertEqual(reqs[0]["evidence_refs"], [])


class DocumentScanTests(unittest.TestCase):
    """F05: inspect actual chart values and package validity."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        self.addCleanup(self.tmp.cleanup)

    def _chart_deck(self, title_xml: str) -> Path:
        return fixtures.pptx(self.dir / "chart.pptx", raw={
            "ppt/slides/slide1.xml": "<a:p><a:r><a:t>정상 본문</a:t></a:r></a:p>",
            "ppt/charts/chart1.xml":
                f"<c:chartSpace>{title_xml}<c:cat><c:strRef><c:strCache>"
                "<c:pt idx=\"0\"><c:v>ABC은행</c:v></c:pt>"
                "<c:pt idx=\"1\"><c:v>기타</c:v></c:pt>"
                "</c:strCache></c:strRef></c:cat></c:chartSpace>"})

    def _names(self) -> Path:
        p = self.dir / "names.txt"
        p.write_text("ABC은행\n", encoding="utf-8")
        return p

    def test_stale_name_in_titled_chart_is_found(self):
        # Reproduce title a:t selecting the run path and skipping category c:v values.
        titled = self._chart_deck("<c:title><c:tx><c:rich><a:p><a:r>"
                                  "<a:t>연도별 실적</a:t></a:r></a:p></c:rich></c:tx></c:title>")
        proc = run(QG, titled, "--names", self._names())
        self.assertEqual(proc.returncode, 1, proc.stdout)
        self.assertIn("ABC은행", proc.stdout)

    def test_stale_name_in_untitled_chart_is_found(self):
        proc = run(QG, self._chart_deck(""), "--names", self._names())
        self.assertEqual(proc.returncode, 1, proc.stdout)

    def test_clean_chart_passes(self):
        p = fixtures.pptx(self.dir / "clean.pptx", raw={
            "ppt/slides/slide1.xml": "<a:p><a:r><a:t>정상 본문</a:t></a:r></a:p>",
            "ppt/charts/chart1.xml":
                "<c:chartSpace><c:cat><c:strCache><c:pt><c:v>2025년</c:v></c:pt>"
                "</c:strCache></c:cat></c:chartSpace>"})
        proc = run(QG, p, "--names", self._names())
        self.assertEqual(proc.returncode, 0, proc.stdout)

    def test_zip_renamed_to_pptx_is_usage_error(self):
        fake = self.dir / "fake.pptx"
        with zipfile.ZipFile(fake, "w") as z:
            z.writestr("readme.txt", "hello")
        proc = run(QG, fake)
        self.assertEqual(proc.returncode, 2)
        self.assertIn("필수 파트 없음", proc.stdout + proc.stderr)


class BoundedVerificationTests(unittest.TestCase):
    """A stalled check or incomplete preview cannot provide passing evidence."""

    def setUp(self):
        from unittest.mock import patch
        self.patch = patch
        sys.path.insert(0, str(DOC / "scripts"))
        import unified_gate
        import deck_check
        self.gate, self.deck = unified_gate, deck_check
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)

    def test_child_timeout_and_launch_failure_are_invalid(self):
        calls = [lambda: self.gate.run_deck_check(Path("candidate.pptx")),
                 lambda: self.gate.run_check_numbers(Path("candidate.pptx"), Path("audit.json")),
                 lambda: self.gate.run_quality_gate(Path("candidate.pptx"), "submission", "ko", None, None)]
        for call in calls:
            for failure in (subprocess.TimeoutExpired("check", 240), OSError("cannot launch")):
                with self.subTest(call=call, failure=type(failure).__name__):
                    with self.patch.object(self.gate.subprocess, "run", side_effect=failure) as runner:
                        code, output = call()
                    self.assertEqual(code, 2)
                    self.assertTrue(output)
                    self.assertEqual(runner.call_args.kwargs["timeout"], 240)

    def test_child_success_and_failure_keep_exit_and_output(self):
        for code in (0, 1, 2):
            with self.subTest(code=code):
                proc = subprocess.CompletedProcess(["check"], code, "report", "diagnostic")
                with self.patch.object(self.gate.subprocess, "run", return_value=proc):
                    actual = self.gate.run_deck_check(Path("candidate.pptx"))
                self.assertEqual(actual, (code, "reportdiagnostic"))

    def test_main_propagates_each_child_error_without_ready_status(self):
        import contextlib
        import io
        doc = make_pptx(self.root / "candidate.pptx")
        data = ready_audit()
        data["render"]["artifact_hash"] = digest_of(doc)
        data["package"]["artifact_hash"] = digest_of(doc)
        audit = self.root / "audit.json"
        audit.write_text(json.dumps(data), encoding="utf-8")
        helpers = ("run_quality_gate", "run_deck_check", "run_check_numbers")
        for failed in helpers:
            with self.subTest(failed=failed), contextlib.ExitStack() as stack:
                for helper in helpers:
                    stack.enter_context(self.patch.object(
                        self.gate, helper, return_value=(2, "timed out") if helper == failed else (0, "ok")))
                output = io.StringIO()
                with contextlib.redirect_stdout(output):
                    code = self.gate.main([str(audit), "--doc", str(doc), "--no-explain"])
                self.assertEqual(code, 2, output.getvalue())
                self.assertIn("INVALID", output.getvalue())
                self.assertNotIn("STATUS: SUBMISSION-READY", output.getvalue())

    def _render(self, *, indices=(1, 2), thumbnail_code=0, thumbnail_error=None,
                probe_error=None, missing_thumbnail_tool=False):
        from types import SimpleNamespace
        png = self.root / "previews"
        png.mkdir(exist_ok=True)
        # A previous run has all pages: none may make an incomplete current run pass.
        for index in (1, 2, 3):
            (png / f"slide-{index}.png").write_bytes(b"old-preview")

        def invoke(cmd, **kwargs):
            self.assertGreater(kwargs.get("timeout", 0), 0)
            if "--convert-to" in cmd:
                out = Path(cmd[cmd.index("--outdir") + 1]) / "candidate.pdf"
                out.write_bytes(b"/Type /Page /Type /Page ")
            elif "--version" in cmd or cmd[0] == "pdfinfo":
                if probe_error:
                    raise probe_error
                text = "Pages: 2" if cmd[0] == "pdfinfo" else "LibreOffice test"
                return subprocess.CompletedProcess(cmd, 0, text, "")
            elif cmd[0] == "pdftoppm":
                if thumbnail_error:
                    raise thumbnail_error
                prefix = Path(cmd[-1])
                for index in indices:
                    (prefix.parent / f"slide-{index}.png").write_bytes(b"new-preview")
                return subprocess.CompletedProcess(cmd, thumbnail_code, b"", b"")
            return subprocess.CompletedProcess(cmd, 0, "", "")

        def which(name):
            return None if missing_thumbnail_tool and name == "pdftoppm" else name

        with self.patch.object(self.deck.shutil, "which", side_effect=which), \
                self.patch.object(self.deck.subprocess, "run", side_effect=invoke), \
                self.patch.object(self.deck, "Presentation", return_value=SimpleNamespace(slides=[1, 2])):
            items, evidence = self.deck.render(self.root / "candidate.pptx", png)
        return items, evidence, png

    def test_success_uses_only_current_previews(self):
        items, evidence, png = self._render()
        self.assertFalse(any(self.deck.is_blocking(i) for i in items), items)
        self.assertEqual(evidence["png"], ["slide-1.png", "slide-2.png"])
        self.assertEqual((png / "slide-1.png").read_bytes(), b"new-preview")
        self.assertEqual((png / "slide-3.png").read_bytes(), b"old-preview")

    def test_thumbnail_nonzero_timeout_and_launch_failure_block(self):
        for kwargs in ({"thumbnail_code": 1},
                       {"thumbnail_error": subprocess.TimeoutExpired("pdftoppm", 240)},
                       {"thumbnail_error": OSError("cannot launch")}):
            with self.subTest(kwargs=kwargs):
                items, evidence, _ = self._render(**kwargs)
                self.assertTrue(any(self.deck.is_blocking(i) for i in items), items)
                self.assertEqual(evidence["png"], [])

    def test_missing_or_wrong_pages_cannot_reuse_stale_images(self):
        for indices in ((), (1,), (1, 3), (1, 2, 3)):
            with self.subTest(indices=indices):
                items, evidence, _ = self._render(indices=indices)
                self.assertTrue(any(self.deck.is_blocking(i) for i in items), items)
                self.assertEqual(evidence["png"], [])

    def test_optional_probe_failures_keep_pdf_page_validation(self):
        for failure in (subprocess.TimeoutExpired("probe", 15), OSError("cannot probe")):
            with self.subTest(failure=type(failure).__name__):
                items, evidence, _ = self._render(probe_error=failure)
                self.assertFalse(any(self.deck.is_blocking(i) for i in items), items)
                self.assertEqual(evidence["pdf_pages"], 2)

    def test_requested_png_without_converter_blocks(self):
        items, evidence, png = self._render(missing_thumbnail_tool=True)
        self.assertTrue(any(self.deck.is_blocking(i) for i in items), items)
        self.assertEqual(evidence["png"], [])
        self.assertTrue((png / "candidate.pdf").is_file())

    def test_output_write_failures_block_without_evidence(self):
        with self.patch.object(self.deck.shutil, "copyfile", side_effect=OSError("cannot write")):
            items, evidence, _ = self._render()
        self.assertTrue(any(self.deck.is_blocking(i) for i in items), items)
        self.assertEqual(evidence["png"], [])
        with self.patch.object(self.deck.shutil, "copy", side_effect=OSError("cannot save PDF")):
            items, evidence, _ = self._render(missing_thumbnail_tool=True)
        self.assertTrue(any(self.deck.is_blocking(i) for i in items), items)
        self.assertEqual(evidence["png"], [])

    def test_render_failure_does_not_emit_verified_evidence(self):
        import contextlib
        import io
        doc = make_pptx(self.root / "candidate.pptx")
        result = self.root / "render.json"
        items, evidence, png = self._render(indices=(1,))
        with self.patch.object(self.deck, "render", return_value=(items, evidence)), \
                contextlib.redirect_stdout(io.StringIO()):
            code = self.deck.main([str(doc), "--png-dir", str(png), "--emit-render", str(result)])
        self.assertEqual(code, 1)
        block = json.loads(result.read_text(encoding="utf-8"))
        self.assertFalse(block["render_succeeded"])
        self.assertFalse(block["verified"])
        self.assertFalse(block["visual_review_approved"])
        self.assertEqual(block["artifact_hash"], digest_of(doc))


if __name__ == "__main__":
    unittest.main()
